"""Parse BOM files (CSV or XML) into BomComponent records.

Mirrors the matching logic used by the original web-app prototype: a BOM row
can cover several reference designators at once (e.g. "R1, R2, R5" as one
line for "10k / 0402"), which is the common shape of BOM exports grouped by
part.
"""

from __future__ import annotations

import csv
import io
import re
import uuid
import xml.etree.ElementTree as ET
from typing import Optional

from .models import BomComponent

DESIGNATOR_KEYS = {"designator", "designators", "reference", "references", "refdes", "ref des", "ref"}
VALUE_KEYS = {"value", "val", "comment"}
FOOTPRINT_KEYS = {"footprint", "package", "pattern"}
DESCRIPTION_KEYS = {"description", "desc"}
MANUFACTURER_KEYS = {"manufacturer", "mfg", "mfr"}
MPN_KEYS = {"mpn", "manufacturer part number", "part number", "part_number", "partnumber"}
QTY_KEYS = {"qty", "quantity", "count"}

# Fuzzy fallbacks for the two columns that matter most for identifying a BOM
# at all and for grouping parts by part number, since real-world exports use
# all sorts of spellings ("Ref Des", "Manufacturer P/N", "MPN1"/"MPN2" for
# multiple approved sources, ...) beyond the exact-match sets above. Matched
# against a fully punctuation/space-stripped, lowercased header name.
_DESIGNATOR_RE = re.compile(r"^(ref(erence)?)?des(ignators?)?$|^ref(erences?)?$")
_MPN_RE = re.compile(r"^(mfr|mfg|manufacturer)?part(number|no|num)?\d*$|^mpn\d*$|^(mfr|mfg|manufacturer)?pn\d*$")

# Real Excel BOM exports often have a title/revision block before the actual
# column header row, and the data isn't always on whichever sheet was active
# when the file was last saved (e.g. a cover/summary sheet) — so several
# rows across every sheet are offered as header candidates rather than
# assuming row 1 of the active sheet.
_HEADER_SCAN_ROWS = 20


def _normalize_key(key: str) -> str:
    key = key.strip().lower()
    key = re.sub(r"[_.]", " ", key)
    key = re.sub(r"\s+", " ", key)
    return key


def _compact(key: str) -> str:
    return re.sub(r"[^a-z0-9]", "", key.lower())


def _find_field(row: dict[str, str], candidates: set[str]) -> Optional[str]:
    normalized = {_normalize_key(k): v for k, v in row.items() if k is not None}
    for candidate in candidates:
        value = normalized.get(candidate)
        if value:
            return value
    return None


def _find_field_fuzzy(row: dict[str, str], pattern: re.Pattern) -> Optional[str]:
    for k, v in row.items():
        if k and v and pattern.match(_compact(k)):
            return v
    return None


def looks_like_designator_header(header: list[str]) -> bool:
    for h in header:
        if not h:
            continue
        if _normalize_key(h) in DESIGNATOR_KEYS or _DESIGNATOR_RE.match(_compact(h)):
            return True
    return False


def split_designators(raw: str) -> list[str]:
    parts = re.split(r"[,;/\s]+", raw.strip())
    return [p for p in parts if p]


def _row_to_component(row: dict[str, str]) -> Optional[BomComponent]:
    designator_raw = _find_field(row, DESIGNATOR_KEYS) or _find_field_fuzzy(row, _DESIGNATOR_RE)
    if not designator_raw:
        return None
    designators = split_designators(designator_raw)
    if not designators:
        return None

    qty_raw = _find_field(row, QTY_KEYS)
    quantity = None
    if qty_raw:
        try:
            quantity = int(float(qty_raw))
        except ValueError:
            quantity = None
    if quantity is None:
        quantity = len(designators)

    return BomComponent(
        id=f"bom_{uuid.uuid4()}",
        designators=designators,
        value=_find_field(row, VALUE_KEYS),
        footprint=_find_field(row, FOOTPRINT_KEYS),
        description=_find_field(row, DESCRIPTION_KEYS),
        manufacturer=_find_field(row, MANUFACTURER_KEYS),
        mpn=_find_field(row, MPN_KEYS) or _find_field_fuzzy(row, _MPN_RE),
        quantity=quantity,
    )


def parse_bom_csv(text: str) -> tuple[list[BomComponent], list[str]]:
    warnings: list[str] = []
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        return [], ["Plik CSV jest pusty lub nie zawiera nagłówka."]

    components: list[BomComponent] = []
    for row in reader:
        clean_row = {(k or "").strip(): (v or "").strip() for k, v in row.items() if k is not None}
        component = _row_to_component(clean_row)
        if component:
            components.append(component)

    if not components:
        warnings.append(
            "Nie znaleziono kolumny z oznaczeniami (Designator/Reference/RefDes). "
            "Sprawdź nagłówki pliku CSV."
        )

    return components, warnings


def _collect_xml_components(node: ET.Element, results: list[dict[str, str]]) -> None:
    flat: dict[str, str] = {}
    looks_like_component = False

    for key, value in node.attrib.items():
        flat[key] = value
        if _normalize_key(key) in DESIGNATOR_KEYS:
            looks_like_component = True

    for child in node:
        text = (child.text or "").strip()
        tag = child.tag.split("}")[-1]  # strip XML namespace
        if text and len(list(child)) == 0:
            flat[tag] = text
            if _normalize_key(tag) in DESIGNATOR_KEYS:
                looks_like_component = True

    if looks_like_component:
        results.append(flat)

    for child in node:
        _collect_xml_components(child, results)


def parse_bom_xml(text: str) -> tuple[list[BomComponent], list[str]]:
    try:
        root = ET.fromstring(text)
    except ET.ParseError as exc:
        return [], [f"Nie udało się sparsować pliku XML: {exc}"]

    rows: list[dict[str, str]] = []
    _collect_xml_components(root, rows)

    components: list[BomComponent] = []
    for row in rows:
        component = _row_to_component(row)
        if component:
            components.append(component)

    warnings: list[str] = []
    if not components:
        warnings.append(
            "Nie znaleziono elementów z polem Designator/Reference w pliku XML. "
            "Format XML tego eksportu BOM może nie być obsługiwany — zalecany jest format CSV."
        )

    return components, warnings


def _iter_candidate_header_rows(workbook, max_rows: int = _HEADER_SCAN_ROWS):
    """Yields (sheet_title, row_index, header) for the first `max_rows` rows
    of every worksheet, active sheet first. Real-world BOM exports often
    have a title/revision block before the actual column header row, and
    the data isn't always on whichever sheet was active when the file was
    last saved (e.g. a cover/summary sheet) — so every plausible row is
    offered to the caller instead of assuming row 1 of the active sheet.
    """
    sheets = [workbook.active] + [ws for ws in workbook.worksheets if ws is not workbook.active]
    for ws in sheets:
        for i, row in enumerate(ws.iter_rows(values_only=True, max_row=max_rows)):
            header = [str(c).strip() if c is not None else "" for c in row]
            if any(header):
                yield ws.title, i, header


def iter_xlsx_header_candidates(path, max_rows: int = _HEADER_SCAN_ROWS):
    """Like `_iter_candidate_header_rows`, but opens the workbook itself —
    used by discovery to sniff a candidate BOM/pick-and-place file without
    duplicating the load/parse logic here."""
    try:
        import openpyxl
    except ImportError:
        return
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return
    yield from _iter_candidate_header_rows(workbook, max_rows=max_rows)


def parse_bom_xlsx(path: str) -> tuple[list[BomComponent], list[str]]:
    try:
        import openpyxl
    except ImportError:
        return [], [
            "Brak biblioteki 'openpyxl', wymaganej do odczytu BOM w formacie Excel (.xlsx). "
            "Zainstaluj ją poleceniem 'pip install openpyxl' albo wyeksportuj BOM do CSV."
        ]

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return [], [f"Nie udało się otworzyć pliku Excel: {exc}"]

    header_info = None
    for sheet_title, row_index, header in _iter_candidate_header_rows(workbook):
        if looks_like_designator_header(header):
            header_info = (sheet_title, row_index, header)
            break

    if header_info is None:
        return [], [
            f"Nie znaleziono kolumny z oznaczeniami (Designator/Reference/RefDes) w żadnym z "
            f"{len(workbook.sheetnames)} arkuszy pliku Excel (przeszukano pierwsze {_HEADER_SCAN_ROWS} "
            "wierszy każdego arkusza). Sprawdź nagłówki albo wyeksportuj właściwy arkusz osobno do CSV."
        ]

    sheet_title, header_row_index, header = header_info
    sheet = workbook[sheet_title]
    components: list[BomComponent] = []
    for raw_row in sheet.iter_rows(min_row=header_row_index + 2, values_only=True):
        row = {
            header[i]: ("" if i >= len(raw_row) or raw_row[i] is None else str(raw_row[i]).strip())
            for i in range(len(header))
        }
        if not any(row.values()):
            continue
        component = _row_to_component(row)
        if component:
            components.append(component)

    warnings: list[str] = []
    if not components:
        warnings.append(
            f"Znaleziono nagłówek z oznaczeniami w arkuszu '{sheet_title}' (wiersz {header_row_index + 1}), "
            "ale nie udało się z niego odczytać żadnych komponentów."
        )
    elif sheet_title != workbook.sheetnames[0]:
        warnings.append(f"BOM odczytano z arkusza '{sheet_title}' (nie pierwszego w pliku).")
    return components, warnings


def parse_bom_file(path: str) -> tuple[list[BomComponent], list[str]]:
    lower = path.lower()
    if lower.endswith(".xlsx"):
        return parse_bom_xlsx(path)
    with open(path, "r", encoding="utf-8-sig") as f:
        text = f.read()
    if lower.endswith(".xml"):
        return parse_bom_xml(text)
    return parse_bom_csv(text)
